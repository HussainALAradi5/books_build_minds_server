"""Reusable, bounded table import and export."""
import csv
from io import BytesIO, StringIO

from flask import send_file
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 500
MAX_EXPORT_ROWS = 5000


def safe_cell(value):
    text = str(value if value is not None else "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def export_table(columns, rows, file_format, filename):
    headers = [heading for heading, _ in columns]
    records = [[safe_cell(row.get(key)) for _, key in columns] for row in rows[:MAX_EXPORT_ROWS]]
    output = BytesIO()
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Library"
        sheet.append(headers)
        for record in records:
            sheet.append(record)
        workbook.save(output)
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif file_format == "pdf":
        document = SimpleDocTemplate(output, pagesize=landscape(A4))
        table = Table([headers, *records], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#166534")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        document.build([table])
        mimetype = "application/pdf"
    elif file_format == "csv":
        text = StringIO()
        writer = csv.writer(text)
        writer.writerow(headers)
        writer.writerows(records)
        output.write(text.getvalue().encode("utf-8-sig"))
        mimetype = "text/csv"
    else:
        raise ValueError("format must be xlsx, pdf, or csv")
    output.seek(0)
    return send_file(output, mimetype=mimetype, as_attachment=True, download_name=f"{filename}.{file_format}")


def import_table(upload, required_columns):
    if not upload or not upload.filename:
        raise ValueError("Choose a CSV or XLSX file")
    payload = upload.read(MAX_IMPORT_BYTES + 1)
    if len(payload) > MAX_IMPORT_BYTES:
        raise ValueError("Import file exceeds 2 MB")
    extension = upload.filename.rsplit(".", 1)[-1].lower()
    if extension == "csv":
        reader = csv.DictReader(StringIO(payload.decode("utf-8-sig")))
        rows = list(reader)
    elif extension == "xlsx":
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(cell or "").strip() for cell in next(iterator, [])]
        rows = [dict(zip(headers, values)) for values in iterator]
        workbook.close()
    else:
        raise ValueError("Import must be a CSV or XLSX file")
    if not rows or len(rows) > MAX_IMPORT_ROWS:
        raise ValueError("Import must contain 1 to 500 rows")
    if not set(required_columns).issubset(rows[0]):
        raise ValueError(f"Required columns: {', '.join(required_columns)}")
    return rows
